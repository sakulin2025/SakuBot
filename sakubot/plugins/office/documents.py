import uuid

from sakubot.plugin import Plugin
from sakubot.office_support import source, destination, save_new, result, bounded, convert, render_pdf


def body_paragraphs(parent):
    yield from parent.paragraphs
    seen = set()
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc not in seen:
                    seen.add(cell._tc)
                    yield from body_paragraphs(cell)


class Extension(Plugin):
    name = 'documents'

    def setup(self):
        for name in ('read', 'create', 'replace', 'convert', 'render'):
            self.function(getattr(self, name), name='document_' + name)

    def read(self, path: str, offset: int = 0, limit: int = 30):
        """按正文顺序读取 DOCX 段落和表格，分页返回样式与表格行列信息。offset 从0开始；不把文字提取当作版式检查。旧 .doc 或 .odt 先 document_convert 为 DOCX。批注、修订和页眉页脚需按 Skill 使用 OOXML 读取，不保证出现在正文提取里。"""
        from docx import Document
        from docx.text.paragraph import Paragraph
        if offset < 0 or not 1 <= limit <= 100:
            raise ValueError('offset 需非负，limit 为1到100')
        doc = Document(source(self, path, {'.docx'}))
        rows = []
        total = 0
        for index, block in enumerate(doc.iter_inner_content()):
            total = index + 1
            if index < offset or index >= offset + limit:
                continue
            if isinstance(block, Paragraph):
                rows.append({'index': index, 'type': 'paragraph', 'style': block.style.name,
                    'text': block.text[:8000], 'truncated': len(block.text) > 8000})
            else:
                rows.append({'index': index, 'type': 'table', 'rows_total': len(block.rows), 'columns_total': len(block.columns),
                    'rows': [[cell.text[:1000] for cell in row.cells[:12]] for row in block.rows[:30]],
                    'truncated': len(block.rows) > 30 or len(block.columns) > 12 or any(len(cell.text) > 1000 for row in block.rows[:30] for cell in row.cells[:12])})
        return bounded({'path': path, 'blocks': rows, 'total_blocks': total, 'next_offset': offset + len(rows) if offset + len(rows) < total else None})

    def create(self, output: str, title: str, blocks: list[dict]):
        """创建带中文字体、语义标题、正文和表格的 DOCX 新文件。blocks 的 type 为 paragraph/heading/bullet/table/page_break；text 是文字，heading 可带 level=1到3，table 用 rows 二维数组。不会覆盖原件；完成后 document_render 预览，不要声称仅保存就验证排版。复杂样式、图片、页眉等按 Skill 用 python-docx 脚本。"""
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        target = destination(self, output, {'.docx'})
        if len(blocks) > 1000:
            raise ValueError('正文块过多')
        doc = Document()
        for section in doc.sections:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.top_margin = section.bottom_margin = Cm(2.2)
            section.left_margin = section.right_margin = Cm(2.3)
        for style in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet'):
            font = doc.styles[style].font
            font.name, font.color.rgb = 'Noto Sans CJK SC', RGBColor(0, 0, 0)
            doc.styles[style].element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK SC')
        doc.styles['Normal'].font.size = Pt(10.5)
        doc.styles['Normal'].paragraph_format.space_after = Pt(7)
        doc.styles['Normal'].paragraph_format.line_spacing = 1.25
        if title:
            doc.add_heading(title, level=0)
        for block in blocks:
            kind = block.get('type', 'paragraph')
            if kind == 'page_break':
                doc.add_page_break()
            elif kind == 'table':
                rows = block.get('rows', [])
                if not rows or max(map(len, rows)) > 12:
                    raise ValueError('表格需非空且最多12列')
                table = doc.add_table(rows=0, cols=max(map(len, rows)))
                table.style = 'Table Grid'
                for row in rows:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = str(value)
                header = OxmlElement('w:tblHeader')
                table.rows[0]._tr.get_or_add_trPr().append(header)
                for cell in table.rows[0].cells:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
            elif kind == 'heading':
                level = block.get('level', 1)
                if level not in (1, 2, 3):
                    raise ValueError('标题层级为1到3')
                doc.add_heading(str(block.get('text', '')), level=level)
            elif kind in ('paragraph', 'bullet'):
                doc.add_paragraph(str(block.get('text', '')), style='List Bullet' if kind == 'bullet' else 'Normal')
            else:
                raise ValueError('未知正文块类型')
        if not title and not blocks:
            raise ValueError('文档不能为空')
        save_new(target, doc.save)
        return result(self, target)

    def replace(self, path: str, output: str, changes: list[dict]):
        """按原文唯一匹配修改 DOCX 正文及表格段落，保存新文件，保留原件。changes 每项为 old/new；每个 old 必须在该步恰好出现一次，允许跨普通文字 run，替换文字继承起始 run 格式。不跨段落，不处理超链接、批注或修订；复杂编辑先读 Skill，用 OOXML，不要整体覆盖 paragraph.text 破坏格式。"""
        from docx import Document
        doc = Document(source(self, path, {'.docx'}))
        target = destination(self, output, {'.docx'})
        if not changes or len(changes) > 100:
            raise ValueError('提供1到100条替换')
        paragraphs = list(body_paragraphs(doc))
        for change in changes:
            old, new = change.get('old', ''), change.get('new', '')
            if not isinstance(old, str) or not old or not isinstance(new, str):
                raise ValueError('替换需提供非空 old 和文字 new')
            matches = [p for p in paragraphs if old in p.text]
            if sum(p.text.count(old) for p in matches) != 1:
                raise ValueError('原文匹配不是唯一一次，请读取最新正文后调整 old')
            paragraph = matches[0]
            if ''.join(run.text for run in paragraph.runs) != paragraph.text or paragraph._p.xpath('.//w:ins | .//w:del | .//w:fldChar'):
                raise ValueError('目标含超链接、修订或字段结构，请按 Skill 使用 OOXML 精确编辑')
            start, end = paragraph.text.index(old), paragraph.text.index(old) + len(old)
            position, inserted = 0, False
            for run in paragraph.runs:
                text = run.text
                left, right = position, position + len(text)
                position = right
                if right <= start or left >= end:
                    continue
                prefix = text[:max(0, start - left)]
                suffix = text[max(0, end - left):] if end < right else ''
                run.text = prefix + (new if not inserted else '') + suffix
                inserted = True
        save_new(target, doc.save)
        return result(self, target)

    async def convert(self, path: str, output: str):
        """用独立 LibreOffice 实例转换 DOC/DOCX/ODT/RTF，输出 PDF、DOCX 或 UTF-8 TXT 新文件。保留原件，不运行宏、不更新外链，不连接 Google Docs。转换可能改变复杂排版，最终用预览检查。"""
        file = source(self, path, {'.doc', '.docx', '.odt', '.rtf'})
        target = destination(self, output, {'.pdf', '.docx', '.txt'})
        format = {'.pdf': 'pdf:writer_pdf_Export', '.docx': 'docx:Office Open XML Text', '.txt': 'txt:Text (encoded):UTF8'}[target.suffix.lower()]
        return await convert(self, file, target, format)

    async def render(self, invocation, path: str, pages: list[int] = None):
        """先把 Word 文档转换成临时 PDF，再生成指定页 PNG 预览。默认第1页，每次最多4页；有图像能力的当前模型会收到前两页的临时视觉观察；没有图像能力则只返回路径与附件，不声称已目视检查。"""
        file = source(self, path, {'.doc', '.docx', '.odt', '.rtf'})
        folder = self.workspace.path('output/previews/' + uuid.uuid4().hex)
        folder.mkdir(parents=True)
        pdf = folder / (file.stem + '.pdf')
        await convert(self, file, pdf, 'pdf:writer_pdf_Export')
        return {'pdf': result(self, pdf), **await render_pdf(self, invocation, pdf, pages if pages is not None else [1])}
