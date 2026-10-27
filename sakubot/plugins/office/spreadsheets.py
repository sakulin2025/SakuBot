import csv
import itertools
import tempfile
import uuid
import zipfile
from datetime import date, datetime, time, timedelta
from pathlib import Path

from sakubot.plugin import Plugin
from sakubot.office_support import source, destination, save_new, result, bounded, convert, render_pdf


def cell_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return {'duration_seconds': value.total_seconds()}
    return value


def formula_value(value):
    if isinstance(value, str):
        return value
    return {'type': type(value).__name__, 'text': getattr(value, 'text', None), 'ref': getattr(value, 'ref', None)}


def assign(cell, value):
    if isinstance(value, dict):
        if set(value) == {'date'}:
            parsed = datetime.fromisoformat(value['date'])
            if parsed.tzinfo is not None:
                raise ValueError('Excel 日期不存时区，请先转换为目标当地时间再去除偏移')
            cell.value, cell.number_format = parsed, 'yyyy-mm-dd hh:mm'
        elif set(value) == {'formula'} and isinstance(value['formula'], str) and value['formula'].startswith('='):
            cell.value = value['formula']
        else:
            raise ValueError('单元格对象只支持 {date:ISO时间} 或 {formula:=公式}')
    elif value is None or type(value) in (str, int, float, bool):
        if isinstance(value, str) and len(value) > 32767:
            raise ValueError('单元格文字超过32767字符')
        cell.value = value
        # Explicit formula objects avoid executing spreadsheet-looking external text.
        if isinstance(value, str) and value.startswith('='):
            cell.data_type = 's'
    else:
        raise ValueError('单元格值类型不支持')


class Extension(Plugin):
    name = 'spreadsheets'

    def setup(self):
        for name in ('read', 'create', 'update', 'recalculate', 'errors', 'render'):
            self.function(getattr(self, name), name='spreadsheet_' + name)

    def read(self, path: str, sheet: str = '', range: str = 'A1:L25'):
        """读取 XLSX/XLSM 的工作表列表、指定区域的值、公式与已有缓存；不触发重算，缓存为空不等于零。支持 UTF-8 CSV/TSV，其他旧格式先 spreadsheet_recalculate 转为 XLSX。range 为单个 A1 范围，每次最多2000格、100行、50列。"""
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
        file = source(self, path, {'.xlsx', '.xlsm', '.csv', '.tsv'})
        min_col, min_row, max_col, max_row = range_boundaries(range)
        if not all(type(n) is int and n >= 1 for n in (min_col, min_row, max_col, max_row)) or max_col < min_col or max_row < min_row or max_col > 16384 or max_row > 1048576 or max_col - min_col >= 50 or max_row - min_row >= 100 or (max_col - min_col + 1) * (max_row - min_row + 1) > 2000:
            raise ValueError('读取范围需为1到100行、1到50列且最多2000格')
        if file.suffix.lower() in ('.csv', '.tsv'):
            with file.open(encoding='utf-8-sig', newline='') as stream:
                sample = stream.read(8192)
                stream.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
                except csv.Error:
                    dialect = 'excel-tab' if file.suffix.lower() == '.tsv' else 'excel'
                rows = [row[min_col - 1:max_col] for row in itertools.islice(csv.reader(stream, dialect), min_row - 1, max_row)]
            return bounded({'path': path, 'range': range, 'rows': rows, 'values_are_text': True})
        formulas = load_workbook(file, read_only=True, data_only=False, keep_links=False)
        cached = None
        try:
            cached = load_workbook(file, read_only=True, data_only=True, keep_links=False)
            name = sheet or formulas.sheetnames[0]
            if name not in formulas.sheetnames:
                raise ValueError('工作表不存在，请使用返回的准确名称')
            ws, values = formulas[name], cached[name]
            rows = []
            options = dict(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
            for row, actual in zip(ws.iter_rows(**options), values.iter_rows(**options)):
                rows.append([{'value': cell_value(value.value), **({'formula': formula_value(cell.value)} if cell.data_type == 'f' else {})} for cell, value in zip(row, actual)])
            return bounded({'path': path, 'sheets': formulas.sheetnames, 'sheet': name, 'range': range,
                'dimensions': {'rows': ws.max_row, 'columns': ws.max_column}, 'rows': rows})
        finally:
            formulas.close()
            if cached:
                cached.close()

    def create(self, output: str, sheets: list[dict]):
        """创建 XLSX 新文件。sheets 每项包含 name、rows 二维数组，可选 header=true、freeze='A2'、widths=[列宽]。值可为数字、布尔、null、文本，日期用 {date:ISO无时区时间}，公式用 {formula:'=SUM(B2:B5)'}；普通字符串即便以=开头也保持文本。默认基础排版，复杂图表、格式、验证规则按 Skill 使用 openpyxl。公式需再调用 spreadsheet_recalculate 才有计算缓存。"""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.workbook.properties import CalcProperties
        target = destination(self, output, {'.xlsx'})
        if not 1 <= len(sheets) <= 30:
            raise ValueError('每次创建1到30个工作表')
        book = Workbook()
        book.remove(book.active)
        names, count = set(), 0
        for spec in sheets:
            name = spec.get('name', '')
            if not isinstance(name, str) or not name or len(name) > 31 or any(c in name for c in '[]:*?/\\') or name.lower() in names:
                raise ValueError('工作表名称无效或重复')
            names.add(name.lower())
            ws = book.create_sheet(name)
            ws.sheet_view.showGridLines = False
            for index, row in enumerate(spec.get('rows', []), 1):
                count += len(row)
                if count > 100000:
                    raise ValueError('单次最多写入100000格，请使用脚本构建大表')
                for col, value in enumerate(row, 1):
                    cell = ws.cell(index, col)
                    assign(cell, value)
                    cell.font = Font(name='Noto Sans CJK SC', size=11)
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            from openpyxl.utils import get_column_letter
            for col in range(1, ws.max_column + 1):
                widths = spec.get('widths', [])
                width = widths[col - 1] if col <= len(widths) else 18
                if not isinstance(width, (int, float)) or not 1 <= width <= 100:
                    raise ValueError('列宽为1到100')
                ws.column_dimensions[get_column_letter(col)].width = width
            if spec.get('header', True):
                ws.row_dimensions[1].height = 26
                for cell in ws[1]:
                    cell.fill = PatternFill('solid', fgColor='20364D')
                    cell.font = Font(name='Noto Sans CJK SC', size=11, bold=True, color='FFFFFF')
            ws.freeze_panes = spec.get('freeze', 'A2')
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation, ws.page_setup.paperSize = 'landscape', ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
            ws.print_options.horizontalCentered = True
        book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
        save_new(target, book.save)
        return {**result(self, target), 'recalculated': False}

    def update(self, path: str, output: str, sheet: str, cells: dict, number_formats: dict = None):
        """修改指定工作表的单元格与数字格式，保存新文件。cells 为坐标到值的字典，如 {B2:12,C2:{formula:'=B2*2'}}；number_formats 可指定 {B2:'0.00%'}。支持 XLSX/XLSM，宏文件保留 VBA 但不执行宏。openpyxl 不保证保留所有绘图或不支持的扩展，复杂原件先读 Skill，勿承诺无损。修改后需重算。"""
        from openpyxl import load_workbook
        from openpyxl.utils.cell import coordinate_to_tuple
        file = source(self, path, {'.xlsx', '.xlsm'})
        target = destination(self, output, {file.suffix.lower()})
        if len(cells) > 10000 or len(number_formats or {}) > 10000:
            raise ValueError('每次最多更新10000格')
        book = load_workbook(file, keep_vba=file.suffix.lower() == '.xlsm')
        try:
            if sheet not in book.sheetnames:
                raise ValueError('工作表不存在')
            ws = book[sheet]
            for address in set(cells) | set(number_formats or {}):
                row, col = coordinate_to_tuple(address)
                if not 1 <= row <= 1048576 or not 1 <= col <= 16384:
                    raise ValueError('单元格坐标越界')
            for address, value in cells.items():
                assign(ws[address], value)
            for address, format in (number_formats or {}).items():
                if not isinstance(format, str) or len(format) > 200:
                    raise ValueError('数字格式无效')
                ws[address].number_format = format
            from openpyxl.workbook.properties import CalcProperties
            book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
            save_new(target, book.save)
        finally:
            book.close()
        return {**result(self, target), 'recalculated': False}

    async def recalculate(self, path: str, output: str):
        """通过 LibreOffice 重新计算并导出 XLSX 新文件，也可将 XLS/ODS 转为 XLSX。原件不变，不运行宏或更新外链，不保证 Excel 专有函数、数据表或插件公式兼容。输出后用 spreadsheet_errors 与 spreadsheet_read 检查实际缓存，不能将缓存缺失视为零。宏工作簿不通过此工具转换，避免丢失 VBA。"""
        file = source(self, path, {'.xlsx', '.xls', '.ods'})
        target = destination(self, output, {'.xlsx'})
        if file.suffix.lower() != '.xlsx':
            return await convert(self, file, target, 'xlsx:Calc MS Excel 2007 XML')
        with tempfile.TemporaryDirectory(prefix='sakubot-recalculate-') as folder:
            dirty = Path(folder) / file.name
            from lxml import etree
            with zipfile.ZipFile(file) as archive, zipfile.ZipFile(dirty, 'w', zipfile.ZIP_DEFLATED) as writer:
                for entry in archive.infolist():
                    data = archive.read(entry.filename)
                    if entry.filename == 'xl/workbook.xml':
                        xml = etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))
                        ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                        calc = xml.find(ns + 'calcPr')
                        if calc is None:
                            calc = etree.SubElement(xml, ns + 'calcPr')
                        for key, value in {'calcMode': 'auto', 'fullCalcOnLoad': '1', 'forceFullCalc': '1', 'calcId': '0'}.items():
                            calc.set(key, value)
                        data = etree.tostring(xml, xml_declaration=True, encoding='UTF-8')
                    writer.writestr(entry, data)
            return await convert(self, dirty, target, 'xlsx:Calc MS Excel 2007 XML')

    def errors(self, path: str, max_cells: int = 100000):
        """扫描 XLSX/XLSM 已保存缓存中的 Excel 错误和公式缺失缓存，最多1000000格，返回扫描是否完整。只读，不负责计算；先对修改过的文件重算。没有错误并不代表公式业务逻辑正确，缺失缓存也不代表零。"""
        from openpyxl import load_workbook
        file = source(self, path, {'.xlsx', '.xlsm'})
        if not 1 <= max_cells <= 1000000:
            raise ValueError('max_cells 为1到1000000')
        formulas = load_workbook(file, read_only=True, data_only=False, keep_links=False)
        cached = None
        scanned, missing, blank_strings, errors, stopped = 0, 0, 0, [], False
        try:
            cached = load_workbook(file, read_only=True, data_only=True, keep_links=False)
            for sheet in formulas:
                for row, values in zip(sheet.iter_rows(), cached[sheet.title].iter_rows()):
                    for formula, cell in zip(row, values):
                        if scanned >= max_cells:
                            stopped = True
                            break
                        scanned += 1
                        if formula.data_type == 'f' and cell.value is None:
                            if cell.data_type in ('s', 'inlineStr'):
                                blank_strings += 1
                            else:
                                missing += 1
                        if cell.data_type == 'e' and len(errors) < 100:
                            errors.append({'sheet': sheet.title, 'cell': cell.coordinate, 'error': str(cell.value)})
                    if stopped:
                        break
                if stopped:
                    break
            return {'scanned_cells': scanned, 'complete': not stopped, 'formula_cache_missing': missing,
                'formula_blank_string_results': blank_strings, 'errors': errors, 'errors_capped': len(errors) == 100}
        finally:
            formulas.close()
            if cached:
                cached.close()

    async def render(self, invocation, path: str, pages: list[int] = None):
        """按工作簿打印设置导出 PDF，再渲染指定页 PNG；原件不变。默认第一页，每次最多4页，有图像能力的模型获得前两页的一次性视觉观察。全表可能跨多页，检查 total_pages 后继续分页。指定区域预览请按 Skill 在临时副本设置 print_area，不能为预览改变交付文件。"""
        file = source(self, path, {'.xlsx', '.xlsm', '.xls', '.ods'})
        folder = self.workspace.path('output/previews/' + uuid.uuid4().hex)
        folder.mkdir(parents=True)
        pdf = folder / (file.stem + '.pdf')
        await convert(self, file, pdf, 'pdf:calc_pdf_Export')
        return {'pdf': result(self, pdf), **await render_pdf(self, invocation, pdf, pages if pages is not None else [1])}
